"""Build the Day Trading & TrendSpider Bible from manual/*.md.

Outputs (dist/):
  Day-Trading-TrendSpider-Bible-v1.1.docx   Word (TOC refreshed via Word if available)
  Day-Trading-TrendSpider-Bible-v1.1.pdf    PDF exported by Word (falls back to Edge from HTML)
  Day-Trading-TrendSpider-Bible-v1.1.html   single-file HTML for tablet/desktop reading

Requires: pypandoc_binary, python-docx. Optional: Microsoft Word (docx TOC + PDF), Microsoft Edge.
Run:  python tools/build.py
"""
import subprocess
import sys
from pathlib import Path

import pypandoc
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "manual"
DIST = ROOT / "dist"
BUILD = ROOT / "build"
NAME = "Day-Trading-TrendSpider-Bible-v1.1"

ORDER = [
    "00-front-matter.md", "01-part1-market-basics.md", "02-part2-candles.md", "03-part3-indicators.md",
    "04a-part4-system-v1.md", "04b-part4-risk-exits.md", "05-part5-trendspider.md",
    "06-part6-paper-execution.md", "07-part7-evaluation.md", "08-part8-reference.md", "09-appendices.md",
]

# style name -> (fill colour, border colour)
BOXES = {
    "Box Definition": ("EEF3FB", "2F6FDB"),
    "Box Why": ("EAF6EF", "1A9E5A"),
    "Box Mistake": ("FDECEC", "D64545"),
    "Box Rule": ("FFF6E5", "E08A1E"),
    "Box TrendSpider": ("F2ECFB", "7B4FD1"),
    "Box Recommendation": ("EEF7F7", "1E8C8C"),
    "Box Warn": ("FFF2B3", "B58900"),
    "Box Screenshot": ("F3F3F3", "666666"),
}

METADATA = """---
title: "Day Trading & TrendSpider Bible"
subtitle: "Version 1.1 — Personal Training Manual, Operating Runbook and Reference Guide for the RB Day Trading System v1.1"
author: "Prepared for Ranabir"
date: "Created 2026-10-28 · Updated 2026-10-29 (v1.1) · TrendSpider, SignalStack and Alpaca documentation checked 2026-10-28"
lang: en-US
toc-title: "Table of Contents"
---
"""


def combined_markdown() -> Path:
    BUILD.mkdir(exist_ok=True)
    parts = [METADATA] + [(SRC / f).read_text(encoding="utf-8") for f in ORDER]
    out = BUILD / "manual.md"
    out.write_text("\n\n".join(parts), encoding="utf-8")
    return out


def _shade_and_border(ppr, fill, border):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "18" if side == "left" else "4")
        e.set(qn("w:space"), "4")
        e.set(qn("w:color"), border)
        bdr.append(e)
    ppr.append(bdr)
    ppr.append(shd)


def _table_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "A6A6A6")
        borders.append(e)
    tblPr.append(borders)
    # shade header row
    if tbl.rows:
        for cell in tbl.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "1F3A5F")
            tcPr.append(shd)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.font.bold = True


def postprocess_docx(path: Path):
    doc = Document(str(path))
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name in ("Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in styles]:
            styles[name].font.name = "Calibri"
            styles[name].font.size = Pt(10.5)
    for lvl, size, colour in ((1, 20, "1F3A5F"), (2, 14, "2F6FDB"), (3, 12, "1F3A5F")):
        st = styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(colour)
        if lvl == 1:
            st.paragraph_format.page_break_before = True
    for name, (fill, border) in BOXES.items():
        if name in [s.name for s in styles]:
            st = styles[name]
            ppr = st.element.get_or_add_pPr()
            _shade_and_border(ppr, fill, border)
            st.paragraph_format.left_indent = Pt(6)
            st.paragraph_format.right_indent = Pt(6)
            st.paragraph_format.space_after = Pt(2)
            st.font.size = Pt(10)
    for tbl in doc.tables:
        _table_borders(tbl)
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for run in p.runs:
                        run.font.size = Pt(9)
    # landscape-free: set narrow margins for more table room
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(54)
    doc.save(str(path))


def word_refresh_and_pdf(docx: Path, pdf: Path) -> bool:
    """Use Word (COM via PowerShell) to update the TOC field and export a PDF."""
    ps = f"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$w.DisplayAlerts = 0
try {{
  $d = $w.Documents.Open('{docx}')
  foreach ($t in $d.TablesOfContents) {{ $t.Update() }}
  $d.Save()
  $d.ExportAsFixedFormat('{pdf}', 17)
  # Word sometimes drops the COM connection right after a large export; the files are already written.
  try {{ $d.Close([ref]0) }} catch {{}}
}} finally {{ try {{ $w.Quit() }} catch {{}} }}
"""
    try:
        if pdf.exists():
            pdf.unlink()
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=False, timeout=900)
        return pdf.exists()
    except Exception as exc:  # noqa: BLE001
        print("Word export failed:", exc)
        return False


def edge_pdf(html: Path, pdf: Path) -> bool:
    for edge in (r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
                 r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"):
        if Path(edge).exists():
            subprocess.run([edge, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                            f"--print-to-pdf={pdf}", html.as_uri()], check=False, timeout=300)
            return pdf.exists()
    return False


def main():
    DIST.mkdir(exist_ok=True)
    md = combined_markdown()
    common = ["--toc", "--toc-depth=2", "--number-sections", f"--resource-path={SRC}"]

    docx = DIST / f"{NAME}.docx"
    pypandoc.convert_file(str(md), "docx", outputfile=str(docx), extra_args=common)
    postprocess_docx(docx)
    print("docx:", docx)

    html = DIST / f"{NAME}.html"
    pypandoc.convert_file(str(md), "html5", outputfile=str(html),
                          extra_args=common + ["--standalone", "--embed-resources",
                                               f"--css={ROOT / 'tools' / 'manual.css'}",
                                               "--metadata", "pagetitle=Day Trading & TrendSpider Bible"])
    print("html:", html)

    pdf = DIST / f"{NAME}.pdf"
    if "--no-word" not in sys.argv and word_refresh_and_pdf(docx, pdf):
        print("pdf (Word):", pdf)
    elif edge_pdf(html, pdf):
        print("pdf (Edge):", pdf)
    else:
        print("PDF not produced: open the .docx in Word and Save As PDF.")


if __name__ == "__main__":
    main()
