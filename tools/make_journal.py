"""Create journal/RB-v1-Trade-Journal.xlsx (RB v1.1: schema from Chapter 39, formulas from Chapters 16-19 and 40)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "journal" / "RB-v1-Trade-Journal.xlsx"
OUT.parent.mkdir(exist_ok=True)
ROWS = 500
HEAD = PatternFill("solid", fgColor="1F3A5F")
CALC = PatternFill("solid", fgColor="EEF3FB")
BOLD_W = Font(bold=True, color="FFFFFF")

wb = Workbook()

# ------------------------------------------------------------------ Settings
s = wb.active
s.title = "Settings"
s["A1"], s["A1"].font = "RB v1.1 settings (edit the yellow cells)", Font(bold=True, size=13)
rows = [("Current equity ($)", 100000), ("Risk per trade (% of equity)", 0.0025),
        ("Investment cap per ticker ($)", 20000), ("Target (R multiple)", 2),
        ("Minimum buffer ($)", 0.01), ("Buffer (% of signal high/low)", 0.0002),
        ("Minimum stop (ATR multiple)", 1.0)]
for i, (k, v) in enumerate(rows, start=3):
    s.cell(i, 1, k)
    c = s.cell(i, 2, v)
    c.fill = PatternFill("solid", fgColor="FFF2B3")
s["B4"].number_format = s["B8"].number_format = "0.00%"
s["B3"].number_format = s["B5"].number_format = "$#,##0.00"
s["A11"], s["B11"] = "Dollar risk per trade ($R, a maximum)", "=B3*B4"
s["B11"].number_format = "$#,##0.00"
s["A12"] = "Hard rules (any fail = do not trade): LONG C>O / SHORT C<O; 09:45 <= time < 15:30 ET; last completed 15m Close vs EMA20;"
s["A13"] = "not stale (10 min); not chased (> 0.25 ATR past entry); price not through the stop; one position/order per ticker."
s.column_dimensions["A"].width = 48
s.column_dimensions["B"].width = 16

# ------------------------------------------------------------------ Journal
j = wb.create_sheet("Journal")
cols = [
    ("Trade ID", 14, None), ("Date", 11, None), ("Entry time", 9, None), ("Ticker", 8, None),
    ("Direction", 9, None), ("Alert time", 9, None), ("Exit time", 9, None),
    ("Signal high", 10, None), ("Signal low", 10, None),
    ("Buffer", 8, "=IF($E{r}=\"\",\"\",MAX(Settings!$B$7,ROUND(Settings!$B$8*IF($E{r}=\"LONG\",H{r},I{r}),2)))"),
    ("Planned entry", 11, "=IF($E{r}=\"\",\"\",IF($E{r}=\"LONG\",H{r}+J{r},I{r}-J{r}))"),
    ("Actual entry (fill)", 11, None),
    ("Stop", 10, "=IF($E{r}=\"\",\"\",IF($E{r}=\"LONG\",K{r}-O{r},K{r}+O{r}))"),
    ("Equity at entry", 12, None),
    ("R (risk/share)", 10,
     "=IF(OR($E{r}=\"\",AP{r}=\"\"),\"\",ROUNDUP(ROUND(MAX(IF($E{r}=\"LONG\",K{r}-I{r},H{r}-K{r}),Settings!$B$9*AP{r}),6),2))"),
    ("$R budget", 10, "=IF($E{r}=\"\",\"\",N{r}*Settings!$B$4)"),
    ("Shares (min of risk & cap)", 11,
     "=IF(O{r}=\"\",\"\",MIN(INT(ROUND(P{r}/O{r},6)),INT(ROUND(Settings!$B$5/K{r},6))))"),
    ("Planned $ risk", 11, "=IF(O{r}=\"\",\"\",Q{r}*O{r})"),
    ("Target (2R)", 10, "=IF(O{r}=\"\",\"\",IF($E{r}=\"LONG\",K{r}+Settings!$B$6*O{r},K{r}-Settings!$B$6*O{r}))"),
    ("Exit price", 10, None), ("Exit reason", 12, None),
    ("P/L $", 10, "=IF(OR($E{r}=\"\",T{r}=\"\"),\"\",IF($E{r}=\"LONG\",(T{r}-L{r})*Q{r},(L{r}-T{r})*Q{r}))"),
    ("P/L % equity", 10, "=IF(V{r}=\"\",\"\",V{r}/N{r})"),
    ("R result", 9, "=IF(V{r}=\"\",\"\",V{r}/R{r})"),
    ("Entry slippage $/sh", 10, "=IF(OR($E{r}=\"\",L{r}=\"\"),\"\",IF($E{r}=\"LONG\",L{r}-K{r},K{r}-L{r}))"),
    ("MFE (R)", 8, None), ("MAE (R)", 8, None), ("R at 30 min", 9, None),
    ("Opportunity score /100", 9, None), ("RVOL", 7, None), ("RSI", 7, None),
    ("VWAP dist (ATR)", 9, None), ("EMA alignment / sep", 14, None), ("15m confirm", 9, None),
    ("Screenshot before", 18, None), ("Screenshot after", 18, None), ("Market regime", 16, None),
    ("News catalyst?", 10, None), ("Followed rules?", 9, None), ("Mistake", 14, None), ("Notes", 40, None),
    ("ATR(14) 5m", 9, None), ("Opportunity tier", 14, None),
]
for ci, (name, w, _) in enumerate(cols, start=1):
    c = j.cell(1, ci, name)
    c.fill, c.font = HEAD, BOLD_W
    c.alignment = Alignment(wrap_text=True, vertical="center")
    j.column_dimensions[get_column_letter(ci)].width = w
j.row_dimensions[1].height = 42
j.freeze_panes = "E2"
money = {"H", "I", "J", "K", "L", "M", "N", "O", "P", "R", "S", "T", "V", "Y", "AP"}
for r in range(2, ROWS + 2):
    for ci, (_, _, f) in enumerate(cols, start=1):
        col = get_column_letter(ci)
        if f:
            c = j.cell(r, ci, f.format(r=r))
            c.fill = CALC
        else:
            c = j.cell(r, ci)
        if col in money:
            c.number_format = "#,##0.00"
        if col == "W":
            c.number_format = "0.00%"
        if col == "X":
            c.number_format = "+0.00;-0.00;0.00"

dv_dir = DataValidation(type="list", formula1='"LONG,SHORT"', allow_blank=True)
dv_exit = DataValidation(type="list", formula1='"A-Stop,B-Target,C-VWAP,C-EMA,D-Time,Manual-Other"', allow_blank=True)
dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
dv_mis = DataValidation(type="list", allow_blank=True,
                        formula1='"none,chased,moved stop,oversized,early exit,late exit,skipped checklist,wrong account,revenge,other"')
for dv in (dv_dir, dv_exit, dv_yn, dv_mis):
    j.add_data_validation(dv)
for dv, col in ((dv_dir, "E"), (dv_exit, "U"), (dv_yn, "AH"), (dv_yn, "AL"), (dv_yn, "AM"), (dv_mis, "AN")):
    dv.add(f"{col}2:{col}{ROWS + 1}")

# ------------------------------------------------------------------ Stats
st = wb.create_sheet("Stats")
R = f"Journal!X2:X{ROWS + 1}"
PL = f"Journal!V2:V{ROWS + 1}"
stats = [
    ("Trades (closed)", f"=COUNT({R})", "0"),
    ("Wins", f'=COUNTIF({R},">0")', "0"),
    ("Losses", f'=COUNTIF({R},"<=0")', "0"),
    ("Win rate", "=IFERROR(B4/B3,0)", "0.0%"),
    ("Loss rate", "=IFERROR(B5/B3,0)", "0.0%"),
    ("Average winner (R)", f'=IFERROR(AVERAGEIF({R},">0"),0)', "0.00"),
    ("Average loser (R)", f'=IFERROR(AVERAGEIF({R},"<=0"),0)', "0.00"),
    ("Payoff ratio", "=IFERROR(B8/ABS(B9),0)", "0.00"),
    ("Expectancy (R/trade)", "=B6*B8-B7*ABS(B9)", "+0.00;-0.00"),
    ("Total R", f"=SUM({R})", "+0.00;-0.00"),
    ("Total P/L ($)", f"=SUM({PL})", "$#,##0.00"),
    ("Profit factor", f'=IFERROR(SUMIF({R},">0")/ABS(SUMIF({R},"<0")),0)', "0.00"),
    ("Rules followed %", f'=IFERROR(COUNTIF(Journal!AM2:AM{ROWS + 1},"Y")/COUNTA(Journal!AM2:AM{ROWS + 1}),0)', "0.0%"),
    ("LONG expectancy (R)", f'=IFERROR(AVERAGEIF(Journal!E2:E{ROWS + 1},"LONG",{R}),0)', "+0.00;-0.00"),
    ("SHORT expectancy (R)", f'=IFERROR(AVERAGEIF(Journal!E2:E{ROWS + 1},"SHORT",{R}),0)', "+0.00;-0.00"),
    ("Average entry slippage ($/sh)", f"=IFERROR(AVERAGE(Journal!Y2:Y{ROWS + 1}),0)", "0.000"),
    ("Max drawdown (R)", f"=IFERROR(MAX(Stats!F2:F{ROWS + 1}),0)", "0.00"),
]
st["A1"], st["A1"].font = "RB v1 performance (Chapter 40)", Font(bold=True, size=13)
for i, (k, f, fmt) in enumerate(stats, start=3):
    st.cell(i, 1, k)
    c = st.cell(i, 2, f)
    c.number_format = fmt
st.column_dimensions["A"].width = 32
st.column_dimensions["B"].width = 14
# helper columns: cumulative R, running peak, drawdown
st["D1"], st["E1"], st["F1"] = "Cum R", "Peak", "Drawdown"
for r in range(2, ROWS + 2):
    st[f"D{r}"] = f'=IF(Journal!X{r}="","",SUM(Journal!$X$2:X{r}))'
    st[f"E{r}"] = f'=IF(D{r}="","",MAX(0,MAX($D$2:D{r})))'
    st[f"F{r}"] = f'=IF(D{r}="","",E{r}-D{r})'
st["A21"] = "Interpretation: see Chapter 40. Fewer than 50 trades = workflow data only."

# ------------------------------------------------------------------ Sizer
z = wb.create_sheet("Sizer")
z["A1"], z["A1"].font = "Quick position sizer, RB v1.1 (Chapter 37)", Font(bold=True, size=13)
items = [("Direction (LONG/SHORT)", "LONG"), ("Signal high", 255.60), ("Signal low", 255.24), ("ATR(14) 5m", 0.37),
         ("Buffer", '=MAX(Settings!B7,ROUND(Settings!B8*IF(B3="LONG",B4,B5),2))'),
         ("Entry (E)", '=IF(B3="LONG",B4+B7,B5-B7)'),
         ("R = MAX(structural, 1 ATR)", '=ROUNDUP(ROUND(MAX(IF(B3="LONG",B8-B5,B4-B8),Settings!B9*B6),6),2)'),
         ("Stop", '=IF(B3="LONG",B8-B9,B8+B9)'),
         ("$R", "=Settings!B11"), ("Shares by risk", "=INT(ROUND(B11/B9,6))"),
         ("Shares by cap", "=INT(ROUND(Settings!B5/B8,6))"), ("QUANTITY", "=MIN(B12,B13)"),
         ("Planned $ loss", "=B14*B9"), ("Target (2R)", '=IF(B3="LONG",B8+Settings!B6*B9,B8-Settings!B6*B9)'),
         ("Position value", "=B14*B8")]
for i, (k, v) in enumerate(items, start=3):
    z.cell(i, 1, k)
    c = z.cell(i, 2, v)
    if i in (3, 4, 5, 6):
        c.fill = PatternFill("solid", fgColor="FFF2B3")
    else:
        c.number_format = "#,##0.00" if i not in (12, 13, 14) else "0"
z["A14"].font = z["B14"].font = Font(bold=True)
z["A19"] = "IBM example: E 255.65, R 0.41, stop 255.24, target 256.47, 78 shares. Check the hard rules first (Chapter 16)."
z.column_dimensions["A"].width = 26
z.column_dimensions["B"].width = 14

wb.save(OUT)
print("wrote", OUT)
